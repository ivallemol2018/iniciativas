import json
import re
import pandas as pd
from pathlib import Path

from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.utils import get_column_letter


# ===============================================
# CONFIGURATION
# ===============================================
READMES_JSON_FILE = "readmes.json"
OUTPUT_EXCEL_FILE = "iniciativas_apis.xlsx"
ENDPOINT_FILE_PATTERN = "operation-mapping_*.xlsx"

FINAL_COLUMNS = [
    "Iniciativa",
    "Ticket",
    "API",
    "Owner",
    "Estilo",
    "Tipo",
    "Exposicion",
    "Repositorio",
    "Metodo",
    "Endpoint",
    "Descripcion del Endpoint"
]


# ===============================================
# UTILITY FUNCTIONS
# ===============================================
def extract_apis_from_readme(content: str):
    rows = re.findall(
        r"^\|([^|]+)\|([^|]+)\|([^|]+)\|([^|]+)\|([^|]+)\|([^|]+)\|$",
        content,
        flags=re.MULTILINE,
    )

    rows = [r for r in rows if not all(c.strip().startswith("-") for c in r)]
    if rows:
        rows = rows[1:]

    apis = []
    for r in rows:
        if "test" in r[0].lower():
            continue
        
        apis.append({
            "API": r[0].strip(),
            "Owner": r[1].strip(),
            "Estilo": r[2].strip(),
            "Tipo": r[3].strip(),
            "Exposicion": r[4].strip(),
            "Repositorio": r[5].strip(),
        })

    return apis


def read_endpoints_excel(folder: Path):
    excel_files = list(folder.glob(ENDPOINT_FILE_PATTERN))
    if not excel_files:
        return []

    df = pd.read_excel(excel_files[0], header=0)

    if df.shape[1] < 7:
        return[]

    endpoints = []
    for _, row in df.iterrows():
        endpoints.append({
            "API_KEY": str(row.iloc[0]).strip().upper(),
            "Metodo": row.iloc[4],
            "Endpoint": row.iloc[5],
            "Descripcion del Endpoint": row.iloc[6]
        })

    return endpoints


# ===============================================
# MAIN PROCESS
# ===============================================
with open(READMES_JSON_FILE, "r", encoding="utf-8") as f:
    readme_data = json.load(f)

rows = []

for readme in readme_data["readmes"]:
    folder = Path(readme["folder"])
    content = readme["content"]

    initiative_match = re.search(r"Codigo de iniciativa:\s(.+)", content)
    ticket_match = re.search(r"Codigo de ticket:\s(.+)", content)

    iniciativa = initiative_match.group(1).strip() if initiative_match else None
    ticket = ticket_match.group(1).strip() if ticket_match else None

    apis = extract_apis_from_readme(content)
    endpoints = read_endpoints_excel(folder)

    for api in apis:
        api_key = api["API"].strip().upper()
        api_endpoints = [e for e in endpoints if e["API_KEY"] == api_key]

        if not api_endpoints:
            rows.append({
                "Iniciativa": iniciativa,
                "Ticket": ticket,
                **api,
                "Metodo": None,
                "Endpoint": None,
                "Descripcion del Endpoint": None
            })
            continue

        for ep in api_endpoints:
            rows.append({
                "Iniciativa": iniciativa,
                "Ticket": ticket,
                **api,
                "Metodo": ep["Metodo"],
                "Endpoint": ep["Endpoint"],
                "Descripcion del Endpoint": ep["Descripcion del Endpoint"]
            })

df = pd.DataFrame(rows)[FINAL_COLUMNS]


# ===============================================
# EXPORT TO EXCEL (EXECUTE STYLE)
# ===============================================
with pd.ExcelWriter(OUTPUT_EXCEL_FILE, engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name="APIs")
    ws = writer.book["APIs"]

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Header style (blue like image)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")

    thin  = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Body cells
    for r in range(2, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).border = border
            ws.cell(row=r, column=c).alignment = Alignment(vertical="center")

    # Autp column width (safe)
    for idx, col_name in enumerate(df.columns, start=1):
        col_letter = get_column_letter(idx)
        max_len = max(
            len(col_name),
            df[col_name].fillna("").astype(str).map(len).max()
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 80)

print(f"Excel generado exitosamente: {OUTPUT_EXCEL_FILE}")