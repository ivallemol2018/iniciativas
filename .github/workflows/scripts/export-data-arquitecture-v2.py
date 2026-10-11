import re
import unicodedata
import pandas as pd
from pathlib import Path

from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.utils import get_column_letter


# ===============================================
# CONFIGURATION
# ===============================================
OUTPUT_EXCEL_FILE = "iniciativas_apis.xlsx"
PRODUCTIVO_FILE = ".github/workflows/data/reporte_api_productivo.xlsx"
SBB_CR_BQ_FILE = ".github/workflows/data/reporte_ssb_cr_bq_interno.xlsx"
BIAN_RECORDS_FILE = ".github/workflows/data/bian_records_mapping_v1.xlsx"

FINAL_COLUMNS = [
    "API",
    "Metodo",
    "Endpoint"
]


# ===============================================
# UTILITY FUNCTIONS
# ===============================================
def normalize_text(value) -> str:
    text = "" if value is None else str(value)
    text = text.strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    return text


def normalize_endpoint(value) -> str:
    text = normalize_text(value)
    # Reemplaza cualquier path parameter (ej. {id}, {clienteId}) por un comodin generico
    text = re.sub(r"\{[^}]*\}", "{}", text)
    return text


def normalize_endpoint_path(value) -> str:
    """Igual que normalize_endpoint, pero descartando los query params."""
    path = "" if value is None else str(value)
    return normalize_endpoint(path.split("?", 1)[0])


def load_productive_apis(path: str) -> pd.DataFrame:
    file = Path(path)
    if not file.exists():
        raise FileNotFoundError(f"No se encontro el reporte productivo: {path}")

    df = pd.read_excel(file)

    missing = [c for c in FINAL_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(f"Columnas faltantes en '{path}': {missing}")

    df = df[FINAL_COLUMNS].dropna(subset=["API"])
    for col in FINAL_COLUMNS:
        df[col] = df[col].astype(str).str.strip()

    df = df.drop_duplicates().reset_index(drop=True)

    # Todo lo que proviene del reporte productivo ya esta en produccion
    df["Validacion Semantica"] = None
    df["Produccion"] = "Si"

    return df


def merge_bian_records(df: pd.DataFrame, path: str) -> pd.DataFrame:
    """Cruza bian_records_mapping con las filas ya registradas: si existe
    (apiName contenido en API, mismo metodo y mismo endpoint sin query params
    y con path params como comodin) marca 'Validacion Semantica' = Si; si no
    existe, inserta la fila con 'Validacion Semantica' = Si y 'Produccion' = No.
    Solo se consideran las filas con 'aplica' = SI."""
    file = Path(path)
    if not file.exists():
        raise FileNotFoundError(f"No se encontro el reporte BIAN: {path}")

    bian_df = pd.read_excel(file)

    required = ["apiName", "method", "endpoint"]
    missing = [c for c in required + ["aplica"] if c not in bian_df.columns]
    if missing:
        raise ValueError(f"Columnas faltantes en '{path}': {missing}")

    bian_df = bian_df[bian_df["aplica"].map(normalize_text) == "si"]

    rows = df.to_dict("records")
    keys = [
        (normalize_text(r["API"]), normalize_text(r["Metodo"]), normalize_endpoint_path(r["Endpoint"]))
        for r in rows
    ]

    for _, bian in bian_df.dropna(subset=required).iterrows():
        api_name = normalize_text(bian["apiName"])
        metodo = normalize_text(bian["method"])
        endpoint = normalize_endpoint_path(bian["endpoint"])

        matches = [
            i for i, (k_api, k_metodo, k_endpoint) in enumerate(keys)
            if api_name in k_api and k_metodo == metodo and k_endpoint == endpoint
        ]

        if matches:
            for i in matches:
                rows[i]["Validacion Semantica"] = "Si"
            continue

        rows.append({
            "API": str(bian["apiName"]).strip(),
            "Metodo": str(bian["method"]).strip(),
            "Endpoint": str(bian["endpoint"]).strip(),
            "Validacion Semantica": "Si",
            "Produccion": "No",
        })
        keys.append((api_name, metodo, endpoint))

    return pd.DataFrame(rows, columns=df.columns)


def load_ssb_cr_bq_df(path: str):
    file = Path(path)
    if not file.exists():
        return pd.DataFrame(columns=["ssb", "tipo", "cr/bqs"])

    return pd.read_excel(file)


def find_ssb_cr_bq_match(api_value, metodo_value, endpoint_value, ssb_df):
    norm_api = normalize_text(api_value)
    norm_metodo = normalize_text(metodo_value)
    norm_endpoint = normalize_endpoint(endpoint_value)

    for _, row in ssb_df.iterrows():
        # Basta que la API registrada (o el apiName BIAN insertado) este
        # contenida en la API del reporte SSB
        row_api = normalize_text(row.get("Api"))
        if not norm_api or not row_api or norm_api not in row_api:
            continue

        if normalize_text(row.get("Metodo")) != norm_metodo:
            continue

        if normalize_endpoint(row.get("endpoint")) != norm_endpoint:
            continue

        return row.get("ssb"), row.get("tipo"), row.get("cr/bqs"), row.get("dominio")

    return None, None, None, None


def add_group_header(ws, dataframe: pd.DataFrame, groups: dict):
    """Arma un header de 2 filas sobre una hoja escrita con startrow=1: cada
    grupo ({titulo: [columnas]}) se combina horizontalmente en la fila 1 sobre
    sus columnas, y las columnas sin grupo se combinan verticalmente (filas 1-2)."""
    grouped = set()
    for title, cols in groups.items():
        first = dataframe.columns.get_loc(cols[0]) + 1
        last = dataframe.columns.get_loc(cols[-1]) + 1
        ws.cell(row=1, column=first).value = title
        ws.merge_cells(start_row=1, start_column=first, end_row=1, end_column=last)
        grouped.update(cols)

    for idx, col_name in enumerate(dataframe.columns, start=1):
        if col_name in grouped:
            continue
        ws.cell(row=1, column=idx).value = col_name
        ws.merge_cells(start_row=1, start_column=idx, end_row=2, end_column=idx)


def style_worksheet(ws, dataframe: pd.DataFrame, header_rows: int = 1):
    """Aplica el mismo estilo (header azul, bordes, autofiltro, ancho de
    columna) a cualquier hoja generada a partir de un DataFrame. 'header_rows'
    indica cuantas filas ocupa el header (el autofiltro va en la ultima)."""
    ws.freeze_panes = f"A{header_rows + 1}"
    ws.auto_filter.ref = f"A{header_rows}:{get_column_letter(ws.max_column)}{ws.max_row}"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for r in range(1, header_rows + 1):
        for cell in ws[r]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in range(header_rows + 1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).border = border
            ws.cell(row=r, column=c).alignment = Alignment(vertical="center")

    for idx, col_name in enumerate(dataframe.columns, start=1):
        col_letter = get_column_letter(idx)
        max_len = max(
            len(col_name),
            dataframe[col_name].fillna("").astype(str).map(len).max()
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 80)


def fill_column(ws, dataframe: pd.DataFrame, col_name: str, color: str, header_rows: int = 1):
    """Pinta el fondo del header de una columna (la fuente blanca en negrita
    se conserva de style_worksheet). Con header de 2 filas, la fila 1 de una
    columna agrupada es la celda del grupo, que toma el mismo color."""
    col_idx = dataframe.columns.get_loc(col_name) + 1
    fill = PatternFill("solid", fgColor=color)
    for r in range(1, header_rows + 1):
        ws.cell(row=r, column=col_idx).fill = fill


def build_ssb_cr_bq_consolidado(ssb_cr_bq_df: pd.DataFrame, df: pd.DataFrame) -> pd.DataFrame:
    """Consolida reporte_ssb_cr_bq_interno por (SBB, Nombre CR / BQ): el
    conteo preescrito son las filas registradas en ese reporte; los conteos de
    produccion y validacion semantica son las filas cruzadas con esa
    combinacion que tienen 'Produccion' = Si y 'Validacion Semantica' = Si."""
    keys = ["SBB", "Nombre CR / BQ"]
    esperado_col = "# Endpoints Preescrito en Arquitectura"

    esperado_df = (
        ssb_cr_bq_df.rename(columns={"ssb": "SBB", "cr/bqs": "Nombre CR / BQ"})
        .dropna(subset=keys)
        .groupby(keys)
        .size()
        .reset_index(name=esperado_col)
    )

    cruzado_df = df.rename(columns={"Nombre CR/BQ": "Nombre CR / BQ"}).dropna(subset=keys)

    consolidado_df = esperado_df
    for flag_col, count_col in [
        ("Produccion", "# Endpoints Produccion"),
        ("Validacion Semantica", "# Endpoints Validacion Semantica"),
    ]:
        count_df = (
            cruzado_df[cruzado_df[flag_col] == "Si"]
            .groupby(keys)
            .size()
            .reset_index(name=count_col)
        )
        consolidado_df = consolidado_df.merge(count_df, on=keys, how="left")
        consolidado_df[count_col] = consolidado_df[count_col].fillna(0).astype(int)

    return consolidado_df.sort_values(keys).reset_index(drop=True)


# ===============================================
# MAIN PROCESS
# ===============================================
df = load_productive_apis(PRODUCTIVO_FILE)
df = merge_bian_records(df, BIAN_RECORDS_FILE)

ssb_cr_bq_df = load_ssb_cr_bq_df(SBB_CR_BQ_FILE)
ssb_matches = df.apply(
    lambda row: find_ssb_cr_bq_match(row["API"], row["Metodo"], row["Endpoint"], ssb_cr_bq_df),
    axis=1
)
df["Dominio"] = ssb_matches.apply(lambda m: m[3])
df["SBB"] = ssb_matches.apply(lambda m: m[0])
df["CR/BQ"] = ssb_matches.apply(lambda m: m[1])
df["Nombre CR/BQ"] = ssb_matches.apply(lambda m: m[2])


ssb_cr_bq_consolidado_df = build_ssb_cr_bq_consolidado(ssb_cr_bq_df, df)


# ===============================================
# EXPORT TO EXCEL (EXECUTE STYLE)
# ===============================================
with pd.ExcelWriter(OUTPUT_EXCEL_FILE, engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name="Consolidado", startrow=1)
    ws_consolidado = writer.book["Consolidado"]
    add_group_header(ws_consolidado, df, {
        "Preescrito en Arquitectura de Solucion": ["Dominio", "SBB", "CR/BQ", "Nombre CR/BQ"],
    })
    style_worksheet(ws_consolidado, df, header_rows=2)
    fill_column(ws_consolidado, df, "Validacion Semantica", "3B1A74", header_rows=2)
    fill_column(ws_consolidado, df, "Produccion", "1F6E24", header_rows=2)
    fill_column(ws_consolidado, df, "Dominio", "F54927", header_rows=2)
    fill_column(ws_consolidado, df, "SBB", "F54927", header_rows=2)
    fill_column(ws_consolidado, df, "CR/BQ", "F54927", header_rows=2)
    fill_column(ws_consolidado, df, "Nombre CR/BQ", "F54927", header_rows=2)

    ssb_cr_bq_consolidado_df.to_excel(writer, index=False, sheet_name="Resumen")
    style_worksheet(writer.book["Resumen"], ssb_cr_bq_consolidado_df)
    fill_column(writer.book["Resumen"], ssb_cr_bq_consolidado_df, "SBB", "F54927")
    fill_column(writer.book["Resumen"], ssb_cr_bq_consolidado_df, "Nombre CR / BQ", "F54927")
    fill_column(writer.book["Resumen"], ssb_cr_bq_consolidado_df, "# Endpoints Preescrito en Arquitectura", "F54927")
    fill_column(writer.book["Resumen"], ssb_cr_bq_consolidado_df, "# Endpoints Validacion Semantica", "3B1A74")
    fill_column(writer.book["Resumen"], ssb_cr_bq_consolidado_df, "# Endpoints Produccion", "1F6E24")

print(f"Excel generado exitosamente: {OUTPUT_EXCEL_FILE}")
